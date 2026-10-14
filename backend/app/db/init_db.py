"""
One-time script to import the Excel student roster into SQLite.

Uses openpyxl with positional column offsets to match exactly how the
notebook (load_db / load_vocab) reads the Excel file.

Usage:
    cd backend
    python -m app.db.init_db
"""
import sys
import re
from pathlib import Path

# Ensure the backend package is importable
sys.path.insert(0, str(Path(__file__).resolve().parent.parent.parent))

import openpyxl
from app.config import EXCEL_ROSTER_PATH, DATABASE_PATH
from app.models.database import Base, engine, SessionLocal, Student, HindiVocab


# ── Constants matching the notebook ─────────────────────
# Sheets where column 0 is a serial number (offset by 1)
HAS_SR = {'MA', 'MSc', 'MTech', 'PhD', 'e-Masters'}

# Keywords that identify a degree string in a row
DEG_KEYS = ['Bachelor', 'Master', 'Doctor', 'Diploma', 'Executive', 'Three']


def _norm(t) -> str:
    """English normalize: collapse whitespace, strip."""
    if t is None:
        return ""
    return re.sub(r'\s+', ' ', str(t)).strip()


def _hn(s) -> str:
    """Hindi normalize (simplified for import)."""
    import unicodedata
    if s is None:
        return ""
    t = unicodedata.normalize("NFC", str(s))
    t = t.replace("\u200c", "").replace("\u200d", "")
    return t.strip()


def _parse_roll(val) -> str | None:
    """Parse a roll number from a cell value, matching notebook logic."""
    if val is None:
        return None
    try:
        return str(int(float(str(val))))
    except (ValueError, TypeError):
        return None


def create_tables():
    """Create all SQLAlchemy tables if they don't exist."""
    print(f"Creating database at: {DATABASE_PATH}")
    Base.metadata.create_all(bind=engine)
    print("✓ Tables created.")


def import_students():
    """
    Read every sheet from the Excel roster and insert students.
    Mirrors the notebook's load_db(path) logic exactly:
      - Uses openpyxl (not pandas)
      - Positional column offsets with HAS_SR serial-number shift
      - Scans row for degree text using DEG_KEYS
    """
    print(f"Reading roster from: {EXCEL_ROSTER_PATH}")

    if not EXCEL_ROSTER_PATH.exists():
        print(f"ERROR: Excel file not found at {EXCEL_ROSTER_PATH}")
        print("       Update EXCEL_ROSTER_PATH in .env or app/config.py")
        return

    wb = openpyxl.load_workbook(str(EXCEL_ROSTER_PATH), data_only=True)
    db = SessionLocal()
    count = 0

    for sn in wb.sheetnames:
        if sn == "Hindi Degree Name":
            continue

        ws = wb[sn]
        off = 1 if sn in HAS_SR else 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            # Roll number
            roll = _parse_roll(row[off] if len(row) > off else None)
            if not roll:
                continue

            # Skip if already exists
            existing = db.query(Student).filter(Student.roll_number == roll).first()
            if existing:
                continue

            # English name
            english_name = _norm(row[off + 1]) if len(row) > off + 1 else ""
            if not english_name or english_name == "None":
                continue

            # Hindi name
            hindi_name = _hn(row[off + 2]) if len(row) > off + 2 else None

            # Programme
            programme = _norm(row[off + 3]) if len(row) > off + 3 else None

            # Degree — scan remaining columns for a string containing DEG_KEYS
            degree = None
            for cell_val in row[off + 3:]:
                if cell_val is None:
                    continue
                s = str(cell_val).strip()
                if any(k in s for k in DEG_KEYS):
                    degree = s
                    break

            # Branch / discipline — try column off+4 or off+5
            branch = None
            for col_idx in [off + 4, off + 5]:
                if len(row) > col_idx and row[col_idx] is not None:
                    candidate = _norm(row[col_idx])
                    if candidate and candidate != degree and not any(k in candidate for k in DEG_KEYS):
                        branch = candidate
                        break

            student = Student(
                roll_number=roll,
                english_name=english_name,
                hindi_name=hindi_name if hindi_name and hindi_name != "None" else None,
                programme=programme if programme and programme != "None" else None,
                degree=degree,
                branch=branch,
            )
            db.add(student)
            count += 1

        print(f"  ✓ Sheet '{sn}' processed")

    db.commit()
    db.close()
    print(f"✓ Imported {count} students total.")


def import_vocab():
    """
    Read the 'Hindi Degree Name' sheet and import vocabulary mappings.
    Mirrors the notebook's load_vocab(wb) logic:
      - Reads rows with r[:6] slicing
      - Two modes: 'deg' (default) and 'minor' (after sentinel)
      - deg mode:  r[1]=English prog, r[2]=discipline, r[4]=Hindi prog, r[5]=Hindi disc
      - minor mode: r[0]=English minor → r[1]=Hindi minor, r[2]=English hon → r[3]=Hindi hon
    """
    if not EXCEL_ROSTER_PATH.exists():
        return

    wb = openpyxl.load_workbook(str(EXCEL_ROSTER_PATH), data_only=True)
    if "Hindi Degree Name" not in wb.sheetnames:
        print("  ⚠ 'Hindi Degree Name' sheet not found — skipping vocab import")
        return

    ws = wb["Hindi Degree Name"]
    db = SessionLocal()
    count = 0
    mode = "deg"

    for row in ws.iter_rows(min_row=2, values_only=True):
        r = list(row[:6]) if len(row) >= 6 else list(row) + [None] * (6 - len(row))

        # Check for mode sentinel
        if r[0] is not None and str(r[0]).strip() == "Minor in English":
            mode = "minor"
            continue

        if mode == "deg":
            # r[1] = English programme (full), r[2] = English discipline
            # r[4] = Hindi programme, r[5] = Hindi discipline
            full = _norm(r[1])
            disc = _norm(r[2])
            hp = str(r[4]).strip() if r[4] else None
            hd = str(r[5]).strip() if r[5] else None

            if full and hp and hp != "None":
                db.add(HindiVocab(category="prog", english_term=full, hindi_term=hp))
                count += 1
            if disc and hd and hd != "None":
                db.add(HindiVocab(category="disc", english_term=disc, hindi_term=hd))
                count += 1

        elif mode == "minor":
            # r[0] = English minor → r[1] = Hindi minor
            # r[2] = English honours → r[3] = Hindi honours
            minor_en = _norm(r[0])
            minor_hi = str(r[1]).strip() if r[1] else None
            hon_en = _norm(r[2])
            hon_hi = str(r[3]).strip() if r[3] else None

            if minor_en and minor_hi and minor_hi != "None":
                db.add(HindiVocab(category="minor", english_term=minor_en, hindi_term=minor_hi))
                count += 1
            if hon_en and hon_hi and hon_hi != "None":
                db.add(HindiVocab(category="hon", english_term=hon_en, hindi_term=hon_hi))
                count += 1

    db.commit()
    db.close()
    print(f"✓ Imported {count} vocabulary entries.")


if __name__ == "__main__":
    print("=" * 50)
    print("  Degree Verification — Database Initializer")
    print("=" * 50)
    create_tables()
    import_students()
    import_vocab()
    print(f"\n✓ Done! Database ready at: {DATABASE_PATH}")
