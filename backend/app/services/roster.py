"""
Student roster & Hindi vocabulary loader — supports both SQLite ORM and Excel (.xlsx).
Extracted from degree_verification_v2_TESSERACT.ipynb (Part 3).
"""
import re
from typing import Optional
import openpyxl
from sqlalchemy.orm import Session

from app.models.database import Student, HindiVocab
from app.services.hindi_utils import norm, hn, key

HAS_SR = {'MA', 'MSc', 'MTech', 'PhD', 'e-Masters'}
DEG_KEYS = ['Bachelor', 'Master', 'Doctor', 'Diploma', 'Executive', 'Three']


def load_vocab(wb: openpyxl.Workbook) -> dict[str, dict[str, str]]:
    """Load Hindi degree/discipline/minor/honours vocabulary from Excel sheet (Colab Part 3)."""
    V: dict[str, dict[str, str]] = dict(prog={}, disc={}, minor={}, hon={})
    if 'Hindi Degree Name' not in wb.sheetnames:
        return V
    mode = 'deg'
    for r in wb['Hindi Degree Name'].iter_rows(min_row=2, values_only=True):
        r = [norm(c) if c is not None else '' for c in r[:6]]
        if not any(r):
            continue
        if r[0] == 'Minor in English':
            mode = 'minor'
            continue
        if mode == 'deg':
            full, disc, hp, hd = r[1], r[2], r[4], r[5]
            if full and hp:
                V['prog'][key(full)] = hp
            if disc and hd:
                for a, b in zip(re.split(r'\s*/\s*', disc), re.split(r'\s*/\s*', hd)):
                    V['disc'][key(a)] = b
        else:
            if r[0] and r[1]:
                V['minor'][key(r[0])] = r[1]
            if r[2] and r[3]:
                V['hon'][key(r[2])] = r[3]
    return V


def load_db(path: str) -> tuple[dict, dict, list]:
    """Load student database and Hindi vocabulary directly from Excel workbook (Colab Part 3)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    db = {}
    conflicts = []
    for sn in wb.sheetnames:
        if sn == 'Hindi Degree Name':
            continue
        off = 1 if sn in HAS_SR else 0
        for row in wb[sn].iter_rows(min_row=2, values_only=True):
            try:
                roll = str(int(float(str(row[off]))))
            except Exception:
                continue
            deg = ''
            for c in row[off + 3:]:
                if c and isinstance(c, str) and any(k in c for k in DEG_KEYS):
                    deg = norm(c)
            rec = dict(
                name=norm(row[off + 1]),
                name_hi=hn(row[off + 2]),
                degree=deg,
                sheet=sn,
                programme=norm(row[off + 3]),
            )
            if roll in db and sn != 'EG_Phase 0' and db[roll]['sheet'] != 'EG_Phase 0':
                conflicts.append((roll, db[roll]['sheet'], sn))
            if roll not in db or sn != 'EG_Phase 0':
                if (
                    roll in db
                    and db[roll]['sheet'] == 'EG_Phase 0'
                    and (db[roll]['name'], db[roll]['degree'], db[roll]['name_hi'])
                    != (rec['name'], rec['degree'], rec['name_hi'])
                ):
                    rec['eg_diff'] = db[roll]
                db[roll] = rec
    return db, load_vocab(wb), conflicts


def get_student_by_roll(roll_number: str, db: Session) -> Optional[dict]:
    """Look up a student record by roll number from SQLite."""
    student = db.query(Student).filter(Student.roll_number == str(roll_number).strip()).first()
    if not student:
        return None
    return {
        "roll_number":  student.roll_number,
        "english_name": norm(student.english_name),
        "hindi_name":   hn(student.hindi_name),
        "name":         norm(student.english_name),
        "name_hi":      hn(student.hindi_name),
        "programme":    norm(student.programme),
        "degree":       norm(student.degree),
        "branch":       norm(student.branch),
        "honours":      norm(student.honours),
        "minors":       norm(student.minors),
    }


def get_vocab(db: Session) -> dict[str, dict[str, str]]:
    """
    Load the Hindi vocabulary dictionary from SQLite, normalized with Colab's `key()` function.
    """
    rows = db.query(HindiVocab).all()
    vocab: dict[str, dict[str, str]] = {"prog": {}, "disc": {}, "minor": {}, "hon": {}}
    for row in rows:
        cat = row.category
        if cat not in vocab:
            vocab[cat] = {}
        vocab[cat][key(row.english_term)] = hn(row.hindi_term)
    return vocab